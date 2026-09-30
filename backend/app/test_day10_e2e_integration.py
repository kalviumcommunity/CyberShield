"""
Day 10 End-to-End Integration and Architectural Verification Test Suite.
Validates the entire pipeline:
Authentication -> Health -> Document Upload (TXT/PDF/DOCX) -> Extraction ->
Chunking -> Embeddings -> FAISS Vector Indexing -> Mitigation Retrieval -> Grounded AI Response
"""

import io
import os
import sys
import docx
try:
    import pymupdf as fitz
except ImportError:
    import fitz

# Ensure backend root is on sys.path
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from fastapi.testclient import TestClient
from app.main import app
from app.database import Base, engine, SessionLocal
from app.models import User, Document, DocumentChunk, Alert, MitigationResult

client = TestClient(app)


def create_test_pdf(text_content: str) -> bytes:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 50), text_content)
    pdf_bytes = doc.write()
    doc.close()
    return pdf_bytes


def create_test_docx(title: str, text_content: str) -> bytes:
    doc = docx.Document()
    doc.add_heading(title, level=1)
    doc.add_paragraph(text_content)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def run_day10_e2e_tests():
    print("=" * 80)
    print("     CYBERSHIELD DAY 10 - FINAL BACKEND E2E INTEGRATION TEST SUITE      ")
    print("=" * 80)

    # 1. Health Endpoint Verification (GET /api/health)
    print("\n[Step 1] Verifying System Health Endpoint (GET /api/health)...")
    health_resp = client.get("/api/health")
    assert health_resp.status_code == 200, f"Health check failed: {health_resp.text}"
    health_data = health_resp.json()
    assert health_data["status"] == "healthy"
    assert health_data["database"] == "connected"
    assert "vector_index" in health_data
    assert "version" in health_data
    assert "timestamp" in health_data
    print(f"  [PASS] Backend Status: {health_data['status']}, Database: {health_data['database']}")

    # Also verify root /health redirect/alias
    root_health_resp = client.get("/health")
    assert root_health_resp.status_code == 200
    assert root_health_resp.json()["status"] == "healthy"
    print("  [PASS] Root /health endpoint alias verified.")

    # 2. Authentication Verification
    print("\n[Step 2] Verifying Authentication & JWT Flows...")
    # Admin login
    admin_login = client.post("/api/auth/login", json={"email": "admin@cybershield.io", "password": "AdminPass123!"})
    assert admin_login.status_code == 200, f"Admin login failed: {admin_login.text}"
    admin_token = admin_login.json()["access_token"]
    admin_headers = {"Authorization": f"Bearer {admin_token}"}
    print("  [PASS] Admin login succeeded. JWT token acquired.")

    # Analyst login
    analyst_login = client.post("/api/auth/login", json={"email": "analyst@cybershield.io", "password": "AnalystPass123!"})
    assert analyst_login.status_code == 200, f"Analyst login failed: {analyst_login.text}"
    analyst_token = analyst_login.json()["access_token"]
    analyst_headers = {"Authorization": f"Bearer {analyst_token}"}
    print("  [PASS] Analyst login succeeded. JWT token acquired.")

    # Profile /me verification
    me_resp = client.get("/api/auth/me", headers=analyst_headers)
    assert me_resp.status_code == 200
    assert me_resp.json()["email"] == "analyst@cybershield.io"
    print("  [PASS] Token authenticated profile inspection verified.")

    # 3. Document Upload & Multi-Format Text Extraction (TXT, PDF, DOCX)
    print("\n[Step 3] Verifying Document Upload & Text Extraction (TXT, PDF, DOCX)...")

    # A. TXT Upload
    txt_text = (
        "Incident Runbook: Linux SSH Brute Force Containment Procedures. "
        "When high volume failed SSH authentication attempts trigger alert SEC-2026-SSH: "
        "1. Immediately block offending remote source IP addresses in iptables and edge firewall. "
        "2. Enforce public-key authentication only and disable PasswordAuthentication in /etc/ssh/sshd_config. "
        "3. Rotate privileged user passwords and inspect /var/log/auth.log for compromised account footholds."
    )
    txt_resp = client.post(
        "/api/documents/upload",
        data={"document_type": "incident_runbook", "title": "Linux SSH Security Runbook"},
        files={"file": ("ssh_runbook.txt", txt_text.encode("utf-8"), "text/plain")},
        headers=admin_headers,
    )
    assert txt_resp.status_code == 201, f"TXT upload failed: {txt_resp.text}"
    txt_id = txt_resp.json()["document_id"]
    print(f"  [PASS] TXT upload & extraction verified. ID: {txt_id}")

    # B. PDF Upload
    pdf_text = (
        "Vulnerability Advisory: OpenSSL Cryptographic Flaw Remediation. "
        "Security advisory CVE-2026-7777 outlines a severe buffer overflow vulnerability. "
        "Mitigation: Upgrade OpenSSL binaries to version 3.2.1 or higher across all load balancers. "
        "Restart affected web service daemons and re-issue compromised TLS session tickets."
    )
    pdf_bytes = create_test_pdf(pdf_text)
    pdf_resp = client.post(
        "/api/documents/upload",
        data={"document_type": "vulnerability_advisory", "title": "OpenSSL Advisory CVE-2026-7777"},
        files={"file": ("openssl_advisory.pdf", pdf_bytes, "application/pdf")},
        headers=admin_headers,
    )
    assert pdf_resp.status_code == 201, f"PDF upload failed: {pdf_resp.text}"
    pdf_id = pdf_resp.json()["document_id"]
    assert "OpenSSL" in pdf_resp.json()["content"]
    print(f"  [PASS] PDF upload & PyMuPDF extraction verified. ID: {pdf_id}")

    # C. DOCX Upload
    docx_text = (
        "Threat Intelligence Report: Operation Crimson Phantom APT Campaign. "
        "Advanced adversary group Crimson Phantom utilizes obfuscated PowerShell payloads and DNS tunneling. "
        "Mitigation: Deploy DNS sinkholes for malicious C2 domains listed in appendix A. "
        "Enable PowerShell Constrained Language Mode and audit Event ID 4104 execution traces."
    )
    docx_bytes = create_test_docx("Operation Crimson Phantom", docx_text)
    docx_resp = client.post(
        "/api/documents/upload",
        data={"document_type": "threat_intelligence", "title": "Crimson Phantom Threat Report"},
        files={"file": ("crimson_phantom.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        headers=admin_headers,
    )
    assert docx_resp.status_code == 201, f"DOCX upload failed: {docx_resp.text}"
    docx_id = docx_resp.json()["document_id"]
    assert "Crimson Phantom" in docx_resp.json()["content"]
    print(f"  [PASS] DOCX upload & python-docx extraction verified. ID: {docx_id}")

    # 4. Text Chunking Verification
    print("\n[Step 4] Verifying Text Chunking (POST /api/documents/{id}/process)...")
    for doc_id, name in [(txt_id, "TXT"), (pdf_id, "PDF"), (docx_id, "DOCX")]:
        proc_resp = client.post(f"/api/documents/{doc_id}/process", headers=admin_headers)
        assert proc_resp.status_code == 200, f"Processing {name} failed: {proc_resp.text}"
        data = proc_resp.json()
        assert data["chunks_created"] > 0
        print(f"  [PASS] Processed {name} (ID: {doc_id}) into {data['chunks_created']} chunk(s).")

    # Verify chunks via GET endpoint
    chunks_resp = client.get(f"/api/documents/{txt_id}/chunks", headers=analyst_headers)
    assert chunks_resp.status_code == 200
    assert len(chunks_resp.json()) > 0
    print(f"  [PASS] Retrieved {len(chunks_resp.json())} stored chunks via API.")

    # 5. Embedding Generation Verification (384 dimensions)
    print("\n[Step 5] Verifying Vector Embedding Generation (POST /api/documents/{id}/embed)...")
    for doc_id, name in [(txt_id, "TXT"), (pdf_id, "PDF"), (docx_id, "DOCX")]:
        embed_resp = client.post(f"/api/documents/{doc_id}/embed", headers=admin_headers)
        assert embed_resp.status_code == 200, f"Embedding {name} failed: {embed_resp.text}"
        data = embed_resp.json()
        assert data["embedding_dimension"] == 384
        assert data["status"] in ("success", "already_embedded")
        print(f"  [PASS] Generated 384-dimensional embeddings for {name} document.")

    # 6. FAISS Semantic Search Verification
    print("\n[Step 6] Verifying FAISS Semantic Vector Search (GET /api/search)...")
    query = "How do I block an SSH brute force attack on Linux?"
    search_resp = client.get(f"/api/search?q={query}&top_k=3", headers=analyst_headers)
    assert search_resp.status_code == 200, f"Search failed: {search_resp.text}"
    results = search_resp.json()
    assert len(results) > 0
    top = results[0]
    assert "chunk_id" in top
    assert "document_title" in top
    assert "similarity_score" in top
    assert top["similarity_score"] > 0.40
    print(f"  [PASS] Search returned top result: '{top['document_title']}' with score {top['similarity_score']:.4f}")

    # 7. Cybersecurity Mitigation Retrieval Verification
    print("\n[Step 7] Verifying Mitigation Retrieval API (POST /api/mitigation/search)...")
    mit_payload = {
        "alert": "Multiple failed SSH root login attempts observed on production bastion host.",
        "top_k": 3,
        "min_threshold": 0.35,
        "severity": "high",
    }
    mit_resp = client.post("/api/mitigation/search", json=mit_payload, headers=analyst_headers)
    assert mit_resp.status_code == 200, f"Mitigation search failed: {mit_resp.text}"
    mit_data = mit_resp.json()
    assert len(mit_data["results"]) > 0
    first_mit = mit_data["results"][0]
    assert "mitigation_text" in first_mit
    assert "relevance_score" in first_mit
    assert first_mit["document_title"] == "Linux SSH Security Runbook"
    print(f"  [PASS] Retrieved mitigation from: '{first_mit['document_title']}' (Score: {first_mit['relevance_score']:.4f})")

    # 8. Grounded AI Response Layer (RAG) & Source Attribution
    print("\n[Step 8] Verifying Grounded AI RAG Answer & Source Attribution (POST /api/mitigation/answer)...")
    rag_payload = {
        "alert": "Multiple failed SSH root login attempts observed on production bastion host.",
        "top_k": 3,
        "min_threshold": 0.35,
        "severity": "high",
    }
    rag_resp = client.post("/api/mitigation/answer", json=rag_payload, headers=analyst_headers)
    assert rag_resp.status_code == 200, f"RAG answer failed: {rag_resp.text}"
    rag_data = rag_resp.json()
    assert rag_data["answer"] != ""
    assert len(rag_data["sources"]) > 0
    # Check that sources attribute properly
    source_titles = [s["document_title"] for s in rag_data["sources"]]
    assert "Linux SSH Security Runbook" in source_titles
    print(f"  [PASS] AI Answer synthesized with verified sources: {source_titles}")

    # Verify fallback behavior for irrelevant query
    irrelevant_resp = client.post(
        "/api/mitigation/answer",
        json={"alert": "What is the capital city of Australia?"},
        headers=analyst_headers,
    )
    assert irrelevant_resp.status_code == 200
    assert irrelevant_resp.json()["answer"] == "Insufficient information found in the available security documents."
    print("  [PASS] Guardrail fallback verified: 'Insufficient information found in the available security documents.'")

    # 9. Re-verify Health Endpoint after Vector Additions
    print("\n[Step 9] Re-verifying System Health after Vector Population...")
    final_health = client.get("/api/health")
    assert final_health.status_code == 200
    h_json = final_health.json()
    assert h_json["vector_index"]["status"] == "ready"
    assert h_json["vector_index"]["indexed_chunks"] > 0
    print(f"  [PASS] FAISS Index Status: {h_json['vector_index']['status']} ({h_json['vector_index']['indexed_chunks']} indexed chunks)")

    print("\n" + "=" * 80)
    print("      ALL DAY 10 E2E BACKEND INTEGRATION TESTS PASSED SUCCESSFULLY!     ")
    print("=" * 80)
    return True


if __name__ == "__main__":
    success = run_day10_e2e_tests()
    if not success:
        sys.exit(1)
