import io
import os
import sys
import docx
try:
    import pymupdf as fitz
except ImportError:
    import fitz  # PyMuPDF fallback
from fastapi.testclient import TestClient

# Add parent directory to sys.path
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from app.main import app
from app.database import Base, engine, SessionLocal
from app.models import Document, DocumentType

client = TestClient(app)


def create_sample_pdf_bytes(text_content: str) -> bytes:
    """
    Generates sample PDF file bytes in memory using PyMuPDF.
    """
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 50), text_content)
    pdf_bytes = doc.write()
    doc.close()
    return pdf_bytes


def create_sample_docx_bytes(title: str, text_content: str) -> bytes:
    """
    Generates sample DOCX file bytes in memory using python-docx.
    """
    doc = docx.Document()
    doc.add_heading(title, level=1)
    doc.add_paragraph(text_content)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def test_document_processing_suite():
    print("\n=== Running Day 3 Document Processing Tests ===")
    
    # Ensure database schema exists
    Base.metadata.create_all(bind=engine)

    # Authenticate as admin
    login_resp = client.post("/api/auth/login", json={"email": "admin@cybershield.io", "password": "AdminPass123!"})
    token = login_resp.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    # 1. Test TXT file upload and text extraction
    print("\n[Test 1] Uploading TXT file (threat_intelligence)...")
    txt_content = "CyberShield Threat Intelligence Advisory: APT29 Spear-phishing campaign detected."
    txt_response = client.post(
        "/api/documents/upload",
        data={"document_type": "threat_intelligence", "title": "APT29 Threat Advisory"},
        files={"file": ("apt29_report.txt", txt_content.encode("utf-8"), "text/plain")},
        headers=headers,
    )
    assert txt_response.status_code == 201, f"TXT Upload failed: {txt_response.text}"
    txt_data = txt_response.json()
    assert txt_data["upload_status"] == "success"
    assert txt_data["filename"] == "apt29_report.txt"
    assert txt_data["document_type"] == "threat_intelligence"
    assert txt_data["extracted_text_length"] == len(txt_content)
    txt_doc_id = txt_data["document_id"]
    print(f"[OK] TXT upload succeeded. Document ID: {txt_doc_id}, extracted length: {txt_data['extracted_text_length']}")

    # 2. Test PDF file upload and PyMuPDF text extraction
    print("\n[Test 2] Uploading PDF file (incident_runbook)...")
    pdf_text = "Incident Runbook: Critical Ransomware Remediation Procedures."
    pdf_bytes = create_sample_pdf_bytes(pdf_text)
    pdf_response = client.post(
        "/api/documents/upload",
        data={"document_type": "incident_runbook", "title": "Ransomware Runbook"},
        files={"file": ("ransomware_runbook.pdf", pdf_bytes, "application/pdf")},
        headers=headers,
    )
    assert pdf_response.status_code == 201, f"PDF Upload failed: {pdf_response.text}"
    pdf_data = pdf_response.json()
    assert pdf_data["upload_status"] == "success"
    assert pdf_data["filename"] == "ransomware_runbook.pdf"
    assert pdf_data["document_type"] == "incident_runbook"
    assert pdf_text in pdf_data["content"]
    pdf_doc_id = pdf_data["document_id"]
    print(f"[OK] PDF upload succeeded. Document ID: {pdf_doc_id}, extracted text: '{pdf_data['content']}'")

    # 3. Test DOCX file upload and python-docx text extraction
    print("\n[Test 3] Uploading DOCX file (vulnerability_advisory)...")
    docx_title = "Vulnerability Advisory CVE-2026-8888"
    docx_text = "Patch advisory for remote code execution in web application gateway."
    docx_bytes = create_sample_docx_bytes(docx_title, docx_text)
    docx_response = client.post(
        "/api/documents/upload",
        data={"document_type": "vulnerability_advisory", "title": "CVE-2026-8888 Advisory"},
        files={"file": ("cve_2026_8888.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        headers=headers,
    )
    assert docx_response.status_code == 201, f"DOCX Upload failed: {docx_response.text}"
    docx_data = docx_response.json()
    assert docx_data["upload_status"] == "success"
    assert docx_data["filename"] == "cve_2026_8888.docx"
    assert docx_data["document_type"] == "vulnerability_advisory"
    assert docx_text in docx_data["content"]
    docx_doc_id = docx_data["document_id"]
    print(f"[OK] DOCX upload succeeded. Document ID: {docx_doc_id}, extracted length: {docx_data['extracted_text_length']}")

    # 4. Test GET /api/documents (List all documents)
    print("\n[Test 4] GET /api/documents...")
    get_all_resp = client.get("/api/documents", headers=headers)
    assert get_all_resp.status_code == 200, f"GET all failed: {get_all_resp.text}"
    docs = get_all_resp.json()
    assert len(docs) >= 3
    print(f"[OK] GET /api/documents returned {len(docs)} documents.")

    # 5. Test GET /api/documents/{document_id}
    print(f"\n[Test 5] GET /api/documents/{pdf_doc_id}...")
    get_doc_resp = client.get(f"/api/documents/{pdf_doc_id}", headers=headers)
    assert get_doc_resp.status_code == 200, f"GET doc by ID failed: {get_doc_resp.text}"
    doc_detail = get_doc_resp.json()
    assert doc_detail["id"] == pdf_doc_id
    assert doc_detail["file_name"] == "ransomware_runbook.pdf"
    assert pdf_text in doc_detail["content"]
    print(f"[OK] GET /api/documents/{pdf_doc_id} verified.")

    # 6. Test Error Handling: Unsupported File Type
    print("\n[Test 6] Error handling for unsupported file type (.exe)...")
    bad_file_resp = client.post(
        "/api/documents/upload",
        data={"document_type": "threat_intelligence"},
        files={"file": ("malicious.exe", b"binary data", "application/octet-stream")},
        headers=headers,
    )
    assert bad_file_resp.status_code == 400
    assert "Unsupported file type" in bad_file_resp.json()["detail"]
    print(f"[OK] Rejected unsupported file type: {bad_file_resp.json()['detail']}")

    # 7. Test Error Handling: Invalid Document Type
    print("\n[Test 7] Error handling for invalid document_type...")
    bad_type_resp = client.post(
        "/api/documents/upload",
        data={"document_type": "invalid_type"},
        files={"file": ("test.txt", b"some content", "text/plain")},
        headers=headers,
    )
    assert bad_type_resp.status_code == 400
    assert "Invalid document_type" in bad_type_resp.json()["detail"]
    print(f"[OK] Rejected invalid document_type: {bad_type_resp.json()['detail']}")

    # 8. Test Error Handling: Non-existent Document ID
    print("\n[Test 8] Error handling for non-existent document ID 999999...")
    not_found_resp = client.get("/api/documents/999999", headers=headers)
    assert not_found_resp.status_code == 404
    print(f"[OK] Returned 404 for non-existent document.")

    print("\n=== All Day 3 Tests Passed Successfully! ===")
    return True


if __name__ == "__main__":
    success = test_document_processing_suite()
    if not success:
        sys.exit(1)
