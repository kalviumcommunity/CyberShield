import os
import uuid
import fitz  # PyMuPDF
import docx
from fastapi import UploadFile

# Directory for storing uploaded files locally
BASE_DIR = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
UPLOAD_DIR = os.path.join(BASE_DIR, "uploads")

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}


def ensure_upload_dir() -> str:
    """
    Ensures that the local uploads directory exists.
    """
    os.makedirs(UPLOAD_DIR, exist_ok=True)
    return UPLOAD_DIR


def get_file_extension(filename: str) -> str:
    """
    Extracts lowercase file extension from filename.
    """
    _, ext = os.path.splitext(filename)
    return ext.lower()


def is_supported_file_type(filename: str) -> bool:
    """
    Checks if the filename has a supported extension (.pdf, .docx, .txt).
    """
    ext = get_file_extension(filename)
    return ext in SUPPORTED_EXTENSIONS


def extract_text_from_pdf(file_path: str) -> str:
    """
    Extracts plain text from a PDF file using PyMuPDF (fitz).
    """
    try:
        doc = fitz.open(file_path)
        extracted = []
        for page in doc:
            text = page.get_text()
            if text and text.strip():
                extracted.append(text.strip())
        doc.close()
        return "\n\n".join(extracted)
    except Exception as e:
        raise ValueError(f"Failed to extract text from PDF file: {str(e)}")


def extract_text_from_docx(file_path: str) -> str:
    """
    Extracts plain text from a DOCX file using python-docx.
    """
    try:
        doc = docx.Document(file_path)
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    paragraphs.append(" | ".join(row_cells))
        return "\n".join(paragraphs)
    except Exception as e:
        raise ValueError(f"Failed to extract text from DOCX file: {str(e)}")


def extract_text_from_txt(file_path: str) -> str:
    """
    Extracts text from a TXT file using standard Python file reading.
    """
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            return f.read()
    except UnicodeDecodeError:
        with open(file_path, "r", encoding="latin-1", errors="ignore") as f:
            return f.read()
    except Exception as e:
        raise ValueError(f"Failed to read TXT file: {str(e)}")


def extract_text(file_path: str, extension: str) -> str:
    """
    Dispatches text extraction based on file extension.
    """
    ext = extension.lower()
    if ext == ".pdf":
        return extract_text_from_pdf(file_path)
    elif ext == ".docx":
        return extract_text_from_docx(file_path)
    elif ext == ".txt":
        return extract_text_from_txt(file_path)
    else:
        raise ValueError(
            f"Unsupported file format '{extension}'. Supported formats are: .pdf, .docx, .txt"
        )


async def save_uploaded_file(file: UploadFile) -> tuple[str, str]:
    """
    Saves an uploaded file to the local uploads directory with a unique prefix.
    Returns a tuple of (saved_file_path, unique_filename).
    """
    upload_dir = ensure_upload_dir()
    original_filename = file.filename or "uploaded_file"
    ext = get_file_extension(original_filename)

    unique_prefix = str(uuid.uuid4())[:8]
    safe_filename = f"{unique_prefix}_{original_filename}"
    saved_file_path = os.path.join(upload_dir, safe_filename)

    contents = await file.read()
    with open(saved_file_path, "wb") as f:
        f.write(contents)

    return saved_file_path, safe_filename
